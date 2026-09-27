"""Source-preserving bulk integration; never relabel secondary data as primary review.

Builds a queryable relational SQLite database and CSV exports from the seven-source
registry and the existing reviewed master. Raw inputs and the primary master are
immutable. Paper links are not independent-trial determinations.
"""
from pathlib import Path
import hashlib
import json
import math
import re
import sqlite3
import unicodedata
import zipfile
import xml.etree.ElementTree as ET

import numpy as np
import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'data/processed/integration_batch1_20260927'
CANDIDATE = ROOT / 'data/enrichment/outputs/additional_meta_v1/expanded_public_effect_registry_NOT_FINAL.csv'
MASTER = ROOT / 'literature/evidence_database.csv'
LU = ROOT / 'data/enrichment/raw/public_effect_sources_v1/lu2020_yield_WUE.xlsx'
LU_REFS = ROOT / 'data/enrichment/raw/public_effect_sources_v1/lu2020_source_references.docx'
CORE = {'rice', 'maize', 'wheat'}
OUTCOMES = {'yield', 'SOC', 'SOC_stock', 'CH4', 'N2O'}


def norm(x):
    return re.sub(r'[^a-z0-9]', '', unicodedata.normalize('NFKD', str(x)).encode('ascii', 'ignore').decode().lower())


def text(x):
    s = str(x).strip()
    return '' if s.lower() in {'nan', 'none', 'n/a', 'null'} else s


def first_value(*values):
    """Coalesce missing fields without dropping real zero climate/coordinates."""
    return next((v for v in values if text(v)), '')


def num(x):
    try:
        n = float(x)
        return n if math.isfinite(n) else np.nan
    except (ValueError, TypeError):
        return np.nan


def truth(x):
    return str(x).strip().lower() in {'true', '1', 'yes'}


def key(prefix, value):
    return prefix + hashlib.sha256(str(value).encode()).hexdigest()[:20]


def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()


def save(df, name):
    df.to_csv(OUT / (name + '.csv'), index=False, encoding='utf-8-sig')


def paper_title_from_ref(ref):
    # Start after author initials, not at the first capitalized phrase in a title.
    m = re.search(r'\.\s+(?=(?:Effects?|Impact|Wheat|Crop|Influence|Durum|Integrative|Evapotranspiration|Yield|Evaluating|In search)\b)', ref)
    if not m:
        return ''
    title = re.split(r'\.\s+|．', ref[m.end():], maxsplit=1)[0]
    return title.strip()


def lu_reference_map(c):
    with zipfile.ZipFile(LU_REFS) as z:
        root = ET.fromstring(z.read('word/document.xml'))
    refs = [''.join(p.itertext()) for p in root.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p')]
    result = {}
    for sid, g in c[c.dataset.eq('Lu2020')].groupby('source_study_id'):
        citation = g.iloc[0].citation
        surname = re.match(r'[A-Za-z]+', citation).group().lower()
        year = re.search(r'(?:19|20)\d{2}', citation).group()
        matches = [r for r in refs if re.match(r'^' + re.escape(surname) + r'(?:,|\s)', r, re.I) and year in r]
        result[sid] = dict(reference_text=matches[0] if len(matches) == 1 else '',
                           title=paper_title_from_ref(matches[0]) if len(matches) == 1 else '',
                           reference_status='unique_surname_year_in_source_references' if len(matches) == 1 else ('ambiguous_source_reference' if matches else 'no_matching_author_year_reference'),
                           reference_candidates=' || '.join(matches))
    return result


def lu_raw_audit():
    w = openpyxl.load_workbook(LU, data_only=True)
    s = w.active
    notes = [str(r[0].value) for r in s.iter_rows() if r[0].value and 'bold characters' in str(r[0].value)]
    assert len(notes) == 1 and 'CV' in notes[0], 'Imputation-font convention changed'
    rows = {}
    for cells in s.iter_rows(min_row=4):
        v = [x.value for x in cells]
        if not isinstance(v[0], (int, float)):
            continue
        rows[cells[0].row] = dict(treatment_mean=num(v[3]), treatment_sd=num(v[4]),
            control_mean=num(v[5]), control_sd=num(v[6]), treatment_n=num(v[13]), control_n=num(v[13]),
            source_lnrr=num(v[8]), source_variance=num(v[9]),
            variance_origin='source_CV_imputed' if cells[4].font.bold or cells[6].font.bold else 'source_reported_or_SE_converted_not_primary_rechecked',
            source_footnote=notes[0], plastic_mulch=text(v[15]), rotation=text(v[36]),
            crop_raw=text(v[18]), location_raw=text(v[38]), country_raw=text(v[39]))
    return rows


def numeric_audit(r):
    mt, mc = r['treatment_mean'], r['control_mean']
    if not (math.isfinite(mt) and math.isfinite(mc) and mt > 0 and mc > 0):
        return np.nan, np.nan, 'missing_or_nonpositive_means', 'not_recomputable'
    lnrr = math.log(mt / mc)
    old = r['source_lnrr']
    state = 'consistent' if math.isfinite(old) and abs(lnrr-old) <= 1e-5 else ('source_missing' if not math.isfinite(old) else 'mismatch')
    st, sc, nt, nc = (r[f] for f in ('treatment_sd', 'control_sd', 'treatment_n', 'control_n'))
    v = np.nan
    if all(math.isfinite(x) for x in (st, sc, nt, nc)) and min(st, sc) >= 0 and min(nt, nc) >= 2 and nt.is_integer() and nc.is_integer():
        v = (st/mt)**2/nt + (sc/mc)**2/nc
    source_v = r['source_variance']
    vs = 'not_recomputable' if not math.isfinite(v) else ('source_missing' if not math.isfinite(source_v) else ('consistent' if math.isclose(v, source_v, rel_tol=1e-4, abs_tol=1e-8) else 'mismatch'))
    return lnrr, v, state, vs


def build():
    OUT.mkdir(parents=True, exist_ok=True)
    inputs = [CANDIDATE, MASTER, LU, LU_REFS]
    hashes = {str(p.relative_to(ROOT)):sha(p) for p in inputs}
    c = pd.read_csv(CANDIDATE, keep_default_na=False, low_memory=False)
    m = pd.read_csv(MASTER, keep_default_na=False)
    strict_ids = set(pd.read_csv(ROOT/'data/processed/stage_20260927/screened_effects.csv').effect_id)
    lr, raw = lu_reference_map(c), lu_raw_audit()
    # Exact title metadata links only; this never declares experiment independence.
    doi_titles = {}
    for _, r in m.iterrows():
        if text(r.paper_title) and text(r.paper_doi):
            doi_titles.setdefault(norm(r.paper_title), set()).add(r.paper_doi.lower())
    rows, corrections = [], []
    for _, r in c.iterrows():
        title = text(r.title)
        ref_status = 'source_title_available' if len(title.split()) >= 6 else 'primary_identity_unresolved'
        ref = text(r.reference_text)
        x = dict(record_id='secondary:' + r.effect_id, source_effect_id=r.effect_id,
            source_dataset=r.dataset, source_doi=r.source_doi, source_sheet=r.source_sheet,
            source_excel_row=num(r.source_excel_row), source_study_id=r.source_study_id,
            source_observation=text(r.source_observation), source_experiment_id=text(r.ExperimentID),
            primary_title=title, primary_reference=ref, primary_identity_status=ref_status,
            primary_doi='', pathway='biochar_return' if r.pathway == 'biochar' else r.pathway,
            crop=text(first_value(r.crop_core,r.crop)).lower().replace('corn','maize'), residue=text(r.residue),
            outcome=r.outcome, unit=text(r.unit), country_raw=text(first_value(r.country,r.Country)),
            country=text(first_value(r.country,r.Country)), location=text(r.location),
            latitude=num(first_value(r.latitude,r.Latitude)), longitude=num(first_value(r.longitude,r.Longitude)),
            duration_raw=text(first_value(r.duration,r.TrialDuration)), duration_unit='days' if r.dataset=='BiocharDS' else 'source_specific_unverified',
            MAT_raw=text(first_value(r.MAT_raw,r.Temperature)), MAP_raw=text(first_value(r.MAP_raw,r.Precipitation)),
            soil_depth=text(r.soil_depth), source_lnrr=num(r.source_lnRR), source_variance=num(r.source_variance),
            source_effect_unverified_scale=num(r.source_effect),
            variance_origin=text(r.variance_provenance) or 'unresolved',
            experiment_id='', independence_status='source_paper_cluster_only_not_independent_trial',
            field_setting_status='unresolved', contrast_status=text(r.contrast_scope) or 'unresolved',
            primary_reviewed=False, strict_primary=False, primary_analysis_eligible=False,
            source_scope_reviewed=False, secondary_descriptive_eligible=False)
        for f in ('treatment_mean','control_mean','treatment_sd','control_sd','treatment_n','control_n'):
            x[f] = num(r[f])
        if r.dataset == 'Lu2020':
            ident = lr[r.source_study_id]
            x.update(primary_title=ident['title'], primary_reference=ident['reference_text'], primary_identity_status=ident['reference_status'])
            a = raw[int(x['source_excel_row'])]
            for f in ('treatment_mean','control_mean','treatment_sd','control_sd','treatment_n','control_n','source_lnrr','source_variance'):
                assert math.isclose(x[f],a[f],rel_tol=1e-10,abs_tol=1e-12), (r.effect_id,f)
            x.update(variance_origin=a['variance_origin'], field_setting_status='field_criterion_documented_in_secondary_methods',
                     contrast_status='residue_vs_no_residue_per_source; arm_specific_management_not_rechecked',
                     source_scope_reviewed=True, duration_unit='years', unit='source_grain_yield_scale_unspecified; within_row_lnratio_only')
            if title != x['primary_title']:
                corrections.append(dict(record_id=x['record_id'],field='primary_title',old=title,new=x['primary_title'],basis=ident['reference_status']))
            # Invalid country values are quarantined, never geocoded from a title.
            if x['country'] not in {'China','Mexico'}:
                corrections.append(dict(record_id=x['record_id'],field='country',old=x['country'],new='',basis='non_country_label_quarantined_pending_primary_verification'))
                x['country']=''
            x['plastic_mulch_raw']=a['plastic_mulch']
        title=x['primary_title']
        dois=doi_titles.get(norm(title),set()) if title else set()
        if len(dois)==1:
            x['primary_doi']=next(iter(dois))
        identity = ('doi:' + x['primary_doi']) if x['primary_doi'] else ('title:' + norm(title) if len(title.split())>=6 else 'unresolved:' + r.source_study_id)
        x['paper_id']=key('paper_',identity)
        x['paper_identity_basis']='exact_doi_or_normalized_title' if not identity.startswith('unresolved:') else 'source_id_only'
        x['observation_id']=key('obs_',r.dataset+'|'+r.source_sheet+'|'+str(r.source_excel_row))
        x['lnrr'],x['variance_recomputed'],x['mean_ratio_check'],x['variance_check']=numeric_audit(x)
        x['percent_change']=math.expm1(x['lnrr'])*100 if math.isfinite(x['lnrr']) else np.nan
        reasons=[]
        if x['crop'] not in CORE: reasons.append('noncore_or_unresolved_crop')
        if x['outcome'] not in OUTCOMES: reasons.append('extension_outcome_or_GHG_boundary_review')
        if x['paper_identity_basis']=='source_id_only': reasons.append('primary_identity_unresolved')
        if truth(r.mapping_quarantined): reasons.append('legacy_identity_quarantine')
        if truth(r.source_exact_duplicate) or truth(r.repeated_view): reasons.append('source_repeat_view')
        if truth(r.possible_repeated_numeric_record): reasons.append('possible_repeated_numeric_values')
        if x['mean_ratio_check'] not in {'consistent','source_missing'}: reasons.append('lnrr_not_verified_from_means')
        if r.dataset=='BiocharDS':
            if r.feedstock_scope != 'field_residue_candidate': reasons.append('feedstock_not_core_verified')
            if r.Treatment != 'B': reasons.append('biochar_cointervention')
        if r.dataset=='Lu2020':
            if not re.search(r'straw|stover|stalk',x['residue'],re.I): reasons.append('residue_identity_unresolved')
            if 'sweet corn' in title.lower(): reasons.append('sweet_corn_yield_basis_review')
            if any(v in title.lower() for v in ('plastic film','film-mulched','double-blank','ridging','tillage')):
                reasons.append('cointervention_contrast_requires_primary_check')
            if not x['country']: reasons.append('country_metadata_conflict')
            if r.source_study_id=='Lu2020_27': reasons.append('burning_removal_comparator_ambiguity')
            x['secondary_descriptive_eligible']=not reasons
        x['integration_issues']=';'.join(reasons)
        x['integration_tier']='SECONDARY_SOURCE_AUDITED_DESCRIPTIVE_ONLY' if x['secondary_descriptive_eligible'] else ('SECONDARY_NUMERIC_PREPARED_NOT_ADMITTED' if not reasons else 'SECONDARY_HELD')
        rows.append(x)
    secondary=pd.DataFrame(rows)
    primary=[]
    for _,r in m.iterrows():
        identity='doi:'+r.paper_doi.lower() if text(r.paper_doi) else 'title:'+norm(r.paper_title)
        x=dict(record_id='primary:'+r.effect_id,source_effect_id=r.effect_id,source_dataset='reviewed_primary_master',
            source_doi=r.paper_doi,primary_doi=r.paper_doi,primary_title=r.paper_title,
            primary_identity_status='existing_primary_review',paper_id=key('paper_',identity),
            paper_identity_basis='existing_primary_identity',experiment_id=r.study_id,
            independence_status=r.independence_resolution,source_study_id=r.study_id,
            observation_id=key('primary_obs_',r.shared_control_group+'|'+r.treatment_arm+'|'+r.experiment_year),
            source_locator=r.source_locator,pathway=r.pathway,crop=r.crop_core,outcome=r.outcome,
            unit=r.outcome_unit,country=r.country,country_raw=r.country,location=r.location,
            primary_reviewed=True,strict_primary=r.effect_id in strict_ids,
            primary_analysis_eligible=r.effect_id in strict_ids,secondary_descriptive_eligible=False,
            integration_tier='PRIMARY_STRICT_DESCRIPTIVE' if r.effect_id in strict_ids else 'PRIMARY_REVIEWED_HELD',
            integration_issues=r.quality_flags,variance_origin=r.variance_provenance,
            lnrr=num(r.lnrr),variance_recomputed=num(r.variance_lnrr),percent_change=num(r.percent_change))
        for f in ('treatment_mean','control_mean','treatment_sd','control_sd','treatment_n','control_n'):
            x[f]=num(r[f])
        primary.append(x)
    primary=pd.DataFrame(primary)
    all_effects=pd.concat([primary,secondary],ignore_index=True,sort=False)
    # Paper-level overlap quarantine is deliberately broader than effect matching.
    main_papers=set(primary.paper_id)
    all_effects['paper_overlaps_primary']=all_effects.paper_id.isin(main_papers) & ~all_effects.primary_reviewed
    all_effects.loc[all_effects.paper_overlaps_primary,'secondary_descriptive_eligible']=False
    all_effects.loc[all_effects.paper_overlaps_primary & all_effects.integration_tier.eq('SECONDARY_SOURCE_AUDITED_DESCRIPTIVE_ONLY'),'integration_tier']='SECONDARY_PRIMARY_PAPER_OVERLAP_HELD'
    aliases=all_effects[['source_dataset','source_study_id','paper_id','primary_title','primary_doi','primary_identity_status']].drop_duplicates()
    groups=aliases.groupby('paper_id').source_dataset.nunique()
    all_effects['cross_source_paper_overlap']=all_effects.paper_id.map(groups).gt(1)
    # Same paper/outcome/numeric effect is a review link, never proof of same experiment.
    all_effects['numeric_signature']=all_effects.apply(lambda r:key('sig_', '|'.join([r.paper_id,r.outcome,str(round(r.lnrr,8)) if pd.notna(r.lnrr) else r.record_id])),axis=1)
    signature_sources=all_effects.groupby('numeric_signature').source_dataset.nunique()
    all_effects['cross_source_numeric_overlap']=all_effects.numeric_signature.map(signature_sources).gt(1)
    # Resolve repeated worksheet views only WITHIN Li2022, using full available
    # context and preserving the maximum within-sheet multiplicity. Never merge
    # merely because effect sizes are equal, and never infer trial independence.
    li=c[c.dataset.eq('Li2022_SOC')].copy()
    sigcols=['source_study_id','outcome','unit','soil_depth','latitude','longitude','duration',
        'MAT_raw','MAP_raw','clay','SOC_initial','soil_pH','tillage','residue','straw_management','fertilizer',
        'control_mean','control_sd','control_n','treatment_mean','treatment_sd','treatment_n','source_lnRR']
    sig=li[sigcols].fillna('').astype(str)
    li['view_context_signature']=pd.util.hash_pandas_object(sig,index=False).astype(str)
    li['within_sheet_occurrence']=li.groupby(['source_sheet','view_context_signature']).cumcount()
    li['canonical_view_key']=li.view_context_signature+'|'+li.within_sheet_occurrence.astype(str)
    li['canonical_effect_id']=li.groupby('canonical_view_key').effect_id.transform('first')
    li['repeated_worksheet_view']=li.effect_id.ne(li.canonical_effect_id)
    li_map=li.set_index('effect_id').canonical_effect_id
    all_effects['canonical_record_id']=all_effects.record_id
    is_li=all_effects.source_dataset.eq('Li2022_SOC')
    all_effects.loc[is_li,'canonical_record_id']='secondary:'+all_effects.loc[is_li,'source_effect_id'].map(li_map)
    all_effects['repeated_worksheet_view']=all_effects.record_id.ne(all_effects.canonical_record_id)
    all_effects['source_locator']=all_effects.apply(lambda r:text(r.source_locator) or f"{r.source_dataset}/{text(r.source_sheet)}!row {r.source_excel_row}",axis=1)
    papers=aliases.groupby('paper_id',sort=True).agg(primary_title=('primary_title','first'),primary_doi=('primary_doi','first'),
        source_datasets=('source_dataset',lambda s:'|'.join(sorted(set(s)))),source_alias_count=('source_study_id','nunique')).reset_index()
    # Primary observation keys remain conservative source-record keys because
    # not every legacy record resolves a unique shared arm/season. Do not create
    # accidental pseudo-pairs from incomplete arm metadata.
    p_mask=all_effects.primary_reviewed
    all_effects.loc[p_mask,'observation_id']=all_effects.loc[p_mask,'record_id'].map(lambda s:key('primary_obs_',s))
    observations=all_effects[['observation_id','source_dataset','source_sheet','source_excel_row','source_study_id','paper_id']].drop_duplicates()
    provenance=all_effects[['record_id','source_effect_id','source_dataset','source_doi','source_locator','integration_tier','integration_issues']]
    summaries=all_effects.groupby(['source_dataset','integration_tier','outcome'],dropna=False).agg(records=('record_id','size'),paper_clusters=('paper_id','nunique')).reset_index()
    lu_df=all_effects[all_effects.source_dataset.eq('Lu2020')].copy()
    descriptive=lu_df[lu_df.secondary_descriptive_eligible].copy()
    # No inferential model/CI: these are source-paper-balanced descriptive checks.
    paper_descriptive=descriptive.groupby(['paper_id','crop']).agg(lnrr=('lnrr','mean'),records=('record_id','size'),
        country=('country','first'),any_source_imputed=('variance_origin',lambda s:s.eq('source_CV_imputed').any())).reset_index()
    if len(paper_descriptive):
        paper_descriptive['percent_change']=np.expm1(paper_descriptive.lnrr)*100
    coverage=descriptive.groupby(['crop','variance_origin']).agg(records=('record_id','size'),paper_clusters=('paper_id','nunique')).reset_index()
    source_rows=[]
    for dataset,g in c.groupby('dataset'):
        meta={}
        if dataset.startswith('Figshare_'):
            meta_path=ROOT/'data/enrichment/raw/additional_meta_v1'/dataset.split('_')[1]/'metadata.json'
            meta=json.loads(meta_path.read_text(encoding='utf-8'))
        source_rows.append(dict(source_dataset=dataset,source_doi=g.source_doi.iloc[0],
            source_title=meta.get('title',''),source_version=meta.get('version',''),
            licence=meta.get('license',{}).get('name','see_original_repository; not_relicensed_here'),
            access_route='reused_public_source',records=len(g),
            local_registry=str(CANDIDATE.relative_to(ROOT)),registry_sha256=sha(CANDIDATE)))
    source_rows.append(dict(source_dataset='reviewed_primary_master',source_doi='',source_title='Project primary evidence master',
        source_version='2026-09-27',licence='per_original_source',access_route='project_processed_data',records=len(m),
        local_registry=str(MASTER.relative_to(ROOT)),registry_sha256=sha(MASTER)))
    # Recheck inherited duplicate flags on the CURRENT Li2022 layer. Earlier
    # flags were computed before redundant worksheets were removed.
    numeric_keys=['source_study_id','outcome','unit','soil_depth','control_mean','control_sd','control_n','treatment_mean','treatment_sd','treatment_n']
    li['current_numeric_duplicate_flag']=li.duplicated(numeric_keys,keep=False)
    tables={'effects_integrated':all_effects,'papers':papers,'source_paper_aliases':aliases,'sources':pd.DataFrame(source_rows),
        'source_observations':observations,'provenance':provenance,'integration_coverage':summaries,
        'lu2020_source_audit':lu_df,'lu2020_secondary_descriptive':descriptive,
        'lu2020_paper_descriptive':paper_descriptive,'lu2020_secondary_coverage':coverage,
        'li2022_view_dedup_audit':li[['effect_id','source_sheet','source_excel_row','source_study_id','outcome','possible_repeated_numeric_record','current_numeric_duplicate_flag','view_context_signature','within_sheet_occurrence','canonical_effect_id','repeated_worksheet_view']],
        'metadata_corrections':pd.DataFrame(corrections)}
    for name,df in tables.items(): save(df,name)
    with sqlite3.connect(OUT/'evidence_integrated.sqlite') as db:
        for name,df in tables.items(): df.to_sql(name,db,if_exists='replace',index=False)
        db.execute('CREATE UNIQUE INDEX IF NOT EXISTS effect_record_key ON effects_integrated(record_id)')
        db.execute('CREATE INDEX IF NOT EXISTS effect_paper_key ON effects_integrated(paper_id)')
        db.execute('CREATE INDEX IF NOT EXISTS effect_endpoint ON effects_integrated(pathway,outcome,crop)')
        db.execute('CREATE VIEW IF NOT EXISTS primary_strict AS SELECT * FROM effects_integrated WHERE primary_analysis_eligible=1')
        db.execute('CREATE VIEW IF NOT EXISTS secondary_descriptive AS SELECT * FROM effects_integrated WHERE secondary_descriptive_eligible=1')
        db.execute('CREATE VIEW IF NOT EXISTS canonical_source_records AS SELECT * FROM effects_integrated WHERE repeated_worksheet_view=0')
        assert db.execute('PRAGMA integrity_check').fetchone()[0]=='ok'
    assert len(all_effects)==len(c)+len(m) and all_effects.record_id.is_unique
    assert int(all_effects.primary_analysis_eligible.sum())==len(strict_ids)
    assert all(sha(p)==hashes[str(p.relative_to(ROOT))] for p in inputs)
    manifest=dict(batch='20260927_batch1',input_sha256=hashes,source_candidate_records=len(c),primary_master_records=len(m),
        integrated_records=len(all_effects),primary_strict_records=len(strict_ids),
        primary_trial_keys=m.study_id.nunique(),new_primary_trials=0,
        paper_identity_clusters=len(papers),paper_identity_clusters_are_not_independent_trials=True,
        lu_source_records=len(lu_df),lu_reference_resolved_rows=int(lu_df.primary_identity_status.eq('unique_surname_year_in_source_references').sum()),
        lu_source_imputed_rows=int(lu_df.variance_origin.eq('source_CV_imputed').sum()),
        lu_secondary_descriptive_records=len(descriptive),lu_secondary_descriptive_paper_clusters=descriptive.paper_id.nunique(),
        lu_numeric_consistent_rows=int(lu_df.mean_ratio_check.eq('consistent').sum()),
        lu_variance_consistent_rows=int(lu_df.variance_check.eq('consistent').sum()),
        cross_source_overlap_paper_clusters=int((groups>1).sum()),
        secondary_rows_on_existing_primary_papers=int(all_effects.paper_overlaps_primary.sum()),
        source_numeric_conflicts=int(secondary.mean_ratio_check.eq('mismatch').sum()),
        li2022_repeated_worksheet_views=int(li.repeated_worksheet_view.sum()),
        li2022_current_numeric_duplicate_flags=int(li.current_numeric_duplicate_flag.sum()),
        li2022_legacy_numeric_duplicate_flags=int(li.possible_repeated_numeric_record.map(truth).sum()),
        canonical_records_after_within_source_view_dedup=int((~all_effects.repeated_worksheet_view).sum()),
        limits=['Primary master unchanged','Secondary rows are not independent trials','No pooled global estimate','Lu source has 153 rows vs reported 146 yield comparisons; unresolved'],
        output_csv_sha256={p.name:sha(p) for p in OUT.glob('*.csv')})
    (OUT/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({k:v for k,v in manifest.items() if k not in {'input_sha256','output_csv_sha256'}},ensure_ascii=False,indent=2))


if __name__=='__main__':
    build()
